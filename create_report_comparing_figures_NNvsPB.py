from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

# Setup
base_dir = r'c:\v31model\O2'

subfolders = [
    "baseline_with_previous_hyperparams_dec25",
    "remove_relu_from_output_dec25",
    "remove_relu_from_output_and_add_L1_at_loss_dec25_lamda_1e-2",
    "remove_relu_from_output_and_add_L1_at_loss_dec25_lamda_1e-4",
    "softplus_output_no_l1_dec25",
    "softplus_output_l1_lamda_1e-4_dec25",
    "elu_output_no_l1_dec25",
    "elu_output_l1_lamda_1e-4_dec25"
]

folder_names = [
    "Baseline",
    "No ReLU",
    "No ReLU + L1 (1e-2)",
    "No ReLU + L1 (1e-4)",
    "Softplus no L1",
    "Softplus (L1)", 
    "ELU (No L1)", 
    "ELU (L1)"
]

folders = [os.path.join(base_dir, f) for f in subfolders]

# Indices used in the previous step 
indices = [327, 57, 12, 379]

document = Document()
document.add_heading('Comparison Report', 0)

# Baseline is always the first one
baseline_folder = folders[0]
baseline_name = folder_names[0]

# Loop over all other folders to compare against baseline
for i in range(1, len(folders)):
    target_folder = folders[i]
    target_name = folder_names[i]
    
    document.add_heading(f'{baseline_name} vs {target_name}', level=1)
    
    # Create a table for side-by-side layout
    table = document.add_table(rows=1, cols=2)
    table.autofit = True # Let it fit the page width
    
    # Header
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = baseline_name
    hdr_cells[1].text = target_name
    
    # Center headers
    for cell in hdr_cells:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    for idx in indices:
        # Add Row for Images
        row_cells = table.add_row().cells
        
        # Image 1 (Baseline)
        img1_path = os.path.join(baseline_folder, f"comparison_plot_sample_{idx}.png")
        if os.path.exists(img1_path):
            paragraph = row_cells[0].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(img1_path, width=Inches(3.0)) 
        else:
            row_cells[0].text = f"Missing image {idx}"
    
        # Image 2 (Target)
        img2_path = os.path.join(target_folder, f"comparison_plot_sample_{idx}.png")
        if os.path.exists(img2_path):
            paragraph = row_cells[1].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(img2_path, width=Inches(3.0))
        else:
            row_cells[1].text = f"Missing image {idx}"
        
        # Add a caption row
        caption_row = table.add_row().cells
        # Merge cells for the caption
        caption_cell = caption_row[0].merge(caption_row[1])
        caption_cell.text = f"Sample Index: {idx}"
        caption_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_page_break()

output_path = r'c:\v31model\comparison_report.docx'
document.save(output_path)
print(f"Report generated at {output_path}")
